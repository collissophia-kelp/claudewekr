"""Export board-ready Word document (.docx) with target list data."""

from datetime import date

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT

from engine.models import Company
from engine.scoring import score_company
from engine.revenue import calculate_revenue
from engine.pipeline import STAGE_ORDER


# ── Colours ──────────────────────────────────────────────────────────

TIER_COLORS = {
    1: RGBColor(0x54, 0x82, 0x35),  # Dark green
    2: RGBColor(0xA9, 0xD0, 0x8E),  # Light green
    3: RGBColor(0xFF, 0xD9, 0x66),  # Amber
    4: RGBColor(0xD9, 0xD9, 0xD9),  # Grey
}
HEADER_BG = RGBColor(0x44, 0x72, 0xC4)


def export_docx(companies: list[Company], config: dict, output_path: str) -> None:
    """Generate a board-ready Word document."""
    doc = Document()

    # Landscape orientation
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    new_width, new_height = section.page_height, section.page_width
    section.page_width = new_width
    section.page_height = new_height
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)

    # Default font
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    active = [c for c in companies if not c.excluded]
    scored = [(c, *score_company(c, config)) for c in active]
    scored.sort(key=lambda x: -x[1])

    # ── Title & Summary ──────────────────────────────────────────

    title = doc.add_heading("Stimblue+ Strategic Target List", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        f"{date.today().strftime('%B %Y')}  |  "
        f"Revenue model: Hectares x Coverage x Application Rate x Applications/Season x Price/L"
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Summary stats
    tier_counts = {1: 0, 2: 0, 3: 0, 4: 0}
    total_rev = 0.0
    total_ha = 0.0
    for c, ws_val, tier, tl in scored:
        tier_counts[tier] = tier_counts.get(tier, 0) + 1
        total_rev += calculate_revenue(c, config)
        total_ha += c.hectares_controlled

    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = summary.add_run(
        f"Total: {len(scored)} companies  |  "
        f"Tier 1: {tier_counts[1]}  |  Tier 2: {tier_counts[2]}  |  "
        f"Tier 3: {tier_counts[3]}  |  Tier 4: {tier_counts[4]}  |  "
        f"Hectares: {total_ha:,.0f}  |  "
        f"Projected Revenue: \u20ac{total_rev:,.0f}"
    )
    run.font.size = Pt(11)
    run.bold = True

    doc.add_paragraph()  # spacer

    # ── Pipeline Funnel ──────────────────────────────────────────

    doc.add_heading("Pipeline Funnel", level=1)

    funnel_headers = ["Stage", "Name", "Count", "Hectares", "Projected Revenue", "Top Companies"]
    funnel_table = _create_table(doc, funnel_headers)

    stage_names = config.get("pipeline_stages", {})
    for stage in reversed(STAGE_ORDER):
        stage_cos = [(c, ws_val) for c, ws_val, _, _ in scored if c.pipeline_stage == stage]
        ha = sum(c.hectares_controlled for c, _ in stage_cos)
        rev = sum(calculate_revenue(c, config) for c, _ in stage_cos)
        top = ", ".join(c.company_name for c, _ in sorted(stage_cos, key=lambda x: -calculate_revenue(x[0], config))[:3])

        _add_row(funnel_table, [
            stage,
            stage_names.get(stage, ""),
            str(len(stage_cos)),
            f"{ha:,.0f}",
            f"\u20ac{rev:,.0f}",
            top,
        ])

    # Total row
    _add_row(funnel_table, [
        "TOTAL", "", str(len(scored)),
        f"{total_ha:,.0f}", f"\u20ac{total_rev:,.0f}", "",
    ], bold=True)

    doc.add_paragraph()

    # ── Tier 1 Targets ───────────────────────────────────────────

    tier1 = [(c, ws, t, tl) for c, ws, t, tl in scored if t == 1]
    if tier1:
        doc.add_heading(f"Tier 1 Targets ({len(tier1)} companies)", level=1)
        _write_company_table(doc, tier1, config)
        doc.add_paragraph()

    # ── Tier 2 Targets ───────────────────────────────────────────

    tier2 = [(c, ws, t, tl) for c, ws, t, tl in scored if t == 2]
    if tier2:
        doc.add_heading(f"Tier 2 Targets ({len(tier2)} companies)", level=1)
        _write_company_table(doc, tier2, config)
        doc.add_paragraph()

    # ── Target List by Pipeline Stage ────────────────────────────

    doc.add_heading("Target List by Pipeline Stage", level=1)

    for stage in reversed(STAGE_ORDER):
        stage_cos = [(c, ws, t, tl) for c, ws, t, tl in scored
                     if c.pipeline_stage == stage and STAGE_ORDER.index(stage) >= 1]
        if not stage_cos:
            continue

        stage_name = stage_names.get(stage, stage)
        doc.add_heading(f"{stage} \u2014 {stage_name} ({len(stage_cos)})", level=2)
        _write_company_table(doc, stage_cos, config)
        doc.add_paragraph()

    # ── Region Summary ───────────────────────────────────────────

    doc.add_heading("Summary by Region", level=1)

    regions: dict[str, dict] = {}
    for c, ws_val, tier, tl in scored:
        r = c.region or "Unknown"
        if r not in regions:
            regions[r] = {"t1": 0, "t2": 0, "t3": 0, "total": 0, "ha": 0.0, "rev": 0.0}
        regions[r]["total"] += 1
        if tier == 1:
            regions[r]["t1"] += 1
        elif tier == 2:
            regions[r]["t2"] += 1
        elif tier == 3:
            regions[r]["t3"] += 1
        regions[r]["ha"] += c.hectares_controlled
        regions[r]["rev"] += calculate_revenue(c, config)

    region_headers = ["Region", "Tier 1", "Tier 2", "Tier 3", "Total", "Hectares", "Revenue"]
    region_table = _create_table(doc, region_headers)

    for region, data in sorted(regions.items(), key=lambda x: -x[1]["rev"]):
        _add_row(region_table, [
            region,
            str(data["t1"]), str(data["t2"]), str(data["t3"]),
            str(data["total"]),
            f"{data['ha']:,.0f}",
            f"\u20ac{data['rev']:,.0f}",
        ])

    doc.add_paragraph()

    # ── Assumptions ──────────────────────────────────────────────

    doc.add_heading("Scoring Weights & Assumptions", level=1)

    weights = config.get("scoring_weights", {})
    weight_headers = ["Metric", "Weight"]
    weight_table = _create_table(doc, weight_headers)
    for key, val in weights.items():
        _add_row(weight_table, [key.replace("_", " ").title(), str(val)])

    doc.add_paragraph()

    thresholds = config.get("tier_thresholds", {})
    thresh_headers = ["Tier", "Minimum Score"]
    thresh_table = _create_table(doc, thresh_headers)
    for key, val in thresholds.items():
        _add_row(thresh_table, [key.replace("_", " ").title(), str(val)])

    doc.add_paragraph()

    rev_assumptions = config.get("revenue_assumptions", {})
    rev_headers = ["Parameter", "Value"]
    rev_table = _create_table(doc, rev_headers)
    for key, val in rev_assumptions.items():
        _add_row(rev_table, [key.replace("_", " ").title(), str(val)])

    # Save
    doc.save(output_path)


# ── Helpers ──────────────────────────────────────────────────────────


def _write_company_table(doc: Document, companies: list, config: dict) -> None:
    """Write a company table with standard columns."""
    headers = ["Company", "HQ", "Region", "Business Type", "Crops",
               "Score", "Tier", "Stage", "Hectares", "Revenue", "Notes"]
    table = _create_table(doc, headers)

    for c, ws_val, tier, tl in companies:
        rev = calculate_revenue(c, config)
        crops = ", ".join(c.main_crops[:3])
        if len(c.main_crops) > 3:
            crops += f" +{len(c.main_crops) - 3}"

        notes = c.notes[:80] + "..." if len(c.notes) > 80 else c.notes

        row_data = [
            c.company_name,
            c.hq_country,
            c.region,
            c.business_type[:30],
            crops,
            f"{ws_val:.2f}",
            tl,
            c.pipeline_stage,
            f"{c.hectares_controlled:,.0f}" if c.hectares_controlled else "—",
            f"\u20ac{rev:,.0f}" if rev else "—",
            notes,
        ]
        row_cells = _add_row(table, row_data)

        # Color the tier cell
        if tier in TIER_COLORS and row_cells:
            _shade_cell(row_cells[6], TIER_COLORS[tier])


def _create_table(doc: Document, headers: list[str]):
    """Create a table with styled headers."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for paragraph in hdr_cells[i].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade_cell(hdr_cells[i], HEADER_BG)

    return table


def _add_row(table, values: list[str], bold: bool = False):
    """Add a data row to a table."""
    row_cells = table.add_row().cells
    for i, val in enumerate(values):
        row_cells[i].text = val
        for paragraph in row_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(8)
                if bold:
                    run.bold = True
    return row_cells


def _shade_cell(cell, color: RGBColor) -> None:
    """Apply background shading to a cell."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), str(color))

    # Access or create tcPr
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    # Remove existing shading
    existing = tcPr.findall(qn("w:shd"))
    for e in existing:
        tcPr.remove(e)
    tcPr.append(shading)
